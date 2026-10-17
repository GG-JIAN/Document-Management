from collections import OrderedDict

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook

from materials.models import Material, MaterialCategory, MaterialFeature


class Command(BaseCommand):
    help = "Import material categories, records, and properties from the supplied workbook."

    def add_arguments(self, parser):
        parser.add_argument("workbook", help="Path to the material-properties .xlsx workbook")
        parser.add_argument(
            "--update-existing", action="store_true",
            help="Replace imported fields for matching category/material records.",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        try:
            workbook = load_workbook(options["workbook"], read_only=False, data_only=True)
        except Exception as exc:
            raise CommandError(f"无法读取 Excel 文件：{exc}") from exc

        index_order = {}
        if "分类目录" in workbook.sheetnames:
            for row_number, row in enumerate(workbook["分类目录"].iter_rows(min_row=1, values_only=True), 1):
                if row_number >= 5 and row and row[0] and not str(row[0]).startswith("合计"):
                    index_order[str(row[0]).strip()] = len(index_order)

        created_materials = 0
        skipped_materials = 0
        feature_count = 0
        sheet_names = [name for name in workbook.sheetnames if name != "分类目录"]
        sheet_names.sort(key=lambda name: index_order.get(name, len(index_order)))

        for category_order, sheet_name in enumerate(sheet_names):
            sheet = workbook[sheet_name]
            category, _ = MaterialCategory.objects.update_or_create(
                name=sheet_name, defaults={"sort_order": category_order}
            )
            headers = self._column_headers(sheet)
            for row_number in range(3, sheet.max_row + 1):
                material_name = self._text(sheet.cell(row_number, 1).value)
                if not material_name:
                    continue
                material = Material.objects.filter(category=category, name=material_name).first()
                existed = material is not None
                if existed and not options["update_existing"]:
                    skipped_materials += 1
                    continue

                notes, references, features = self._read_material_row(sheet, row_number, headers)
                if material is None:
                    material = Material.objects.create(
                        category=category,
                        name=material_name,
                        notes=notes,
                        references=references,
                        source_row=row_number,
                    )
                    created_materials += 1
                else:
                    material.notes = notes
                    material.references = references
                    material.source_row = row_number
                    material.save(update_fields=("notes", "references", "source_row", "updated_at"))
                    material.features.all().delete()

                MaterialFeature.objects.bulk_create([
                    MaterialFeature(
                        material=material,
                        name=name,
                        condition=condition,
                        result=result,
                        sort_order=order,
                    )
                    for order, (name, condition, result) in enumerate(features)
                ])
                feature_count += len(features)

        workbook.close()
        self.stdout.write(self.style.SUCCESS(
            f"导入完成：新增 {created_materials} 种材料，跳过已有 {skipped_materials} 种，"
            f"写入 {feature_count} 条特性数据。"
        ))

    @staticmethod
    def _column_headers(sheet):
        headers = {}
        for column in range(2, sheet.max_column + 1):
            value = sheet.cell(1, column).value
            if value:
                headers[column] = str(value).strip()
        for merged in sheet.merged_cells.ranges:
            if merged.min_row <= 1 <= merged.max_row:
                title = sheet.cell(1, merged.min_col).value
                if title:
                    for column in range(max(2, merged.min_col), merged.max_col + 1):
                        headers[column] = str(title).strip()
        return headers

    @classmethod
    def _read_material_row(cls, sheet, row_number, headers):
        feature_values = OrderedDict()
        notes = []
        references = []
        for column, title in headers.items():
            value = cls._text(sheet.cell(row_number, column).value)
            if not value:
                continue
            lowered = title.casefold()
            if "备注" in title:
                notes.append(value)
                continue
            if "引用" in title or "参考文献" in title:
                references.append(value)
                continue

            subheader = cls._text(sheet.cell(2, column).value)
            if subheader not in ("实验条件", "实验结果"):
                feature_values.setdefault(title, {"condition": [], "result": []})["result"].append(value)
                continue
            key = "condition" if subheader == "实验条件" else "result"
            feature_values.setdefault(title, {"condition": [], "result": []})[key].append(value)

        features = [
            (name, "\n".join(values["condition"]), "\n".join(values["result"]))
            for name, values in feature_values.items()
        ]
        return "\n".join(notes), "\n".join(references), features

    @staticmethod
    def _text(value):
        if value is None:
            return ""
        return str(value).strip()
