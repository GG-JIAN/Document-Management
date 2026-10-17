from io import BytesIO
from urllib.parse import quote

from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Count, Q
from django.http import Http404, HttpResponse
from django.shortcuts import get_object_or_404, render

from accounts.permissions import can_manage_records
from .models import Material, MaterialCategory


@login_required
def material_list(request):
    query = request.GET.get("q", "").strip()
    selected_category = request.GET.get("category", "").strip()
    materials = Material.objects.select_related("category").prefetch_related("features")
    if query:
        materials = materials.filter(
            Q(name__icontains=query)
            | Q(category__name__icontains=query)
            | Q(notes__icontains=query)
            | Q(references__icontains=query)
            | Q(features__name__icontains=query)
            | Q(features__condition__icontains=query)
            | Q(features__result__icontains=query)
        ).distinct()
    if selected_category.isdigit():
        materials = materials.filter(category_id=int(selected_category))
    elif selected_category:
        selected_category = ""

    page_obj = Paginator(materials.order_by("category__sort_order", "name"), 20).get_page(
        request.GET.get("page")
    )
    categories = MaterialCategory.objects.annotate(material_count=Count("materials"))
    return render(request, "materials/material_list.html", {
        "page_obj": page_obj,
        "categories": categories,
        "query": query,
        "selected_category": selected_category,
        "can_manage": can_manage_records(request.user),
    })


@login_required
def material_detail(request, pk):
    material = get_object_or_404(
        Material.objects.select_related("category").prefetch_related("features"), pk=pk
    )
    return render(request, "materials/material_detail.html", {
        "material": material,
        "can_manage": can_manage_records(request.user),
    })


@login_required
def compare_materials(request):
    ids = request.POST.getlist("material_ids") if request.method == "POST" else request.GET.getlist("id")
    valid_ids = []
    for value in ids:
        if value.isdigit() and int(value) not in valid_ids:
            valid_ids.append(int(value))
    if len(valid_ids) < 2:
        return render(request, "materials/material_compare.html", {
            "materials": [], "comparison_rows": [], "error": "请选择至少两种材料进行对比。",
        })
    if len(valid_ids) > 6:
        return render(request, "materials/material_compare.html", {
            "materials": [], "comparison_rows": [], "error": "一次最多对比 6 种材料。",
        })

    found = {
        material.pk: material
        for material in Material.objects.select_related("category").prefetch_related("features").filter(pk__in=valid_ids)
    }
    if len(found) != len(valid_ids):
        raise Http404("部分材料已不存在，请返回列表重新选择。")
    materials = [found[pk] for pk in valid_ids]

    feature_names = {}
    for material in materials:
        for feature in material.features.all():
            feature_names[feature.name] = min(feature_names.get(feature.name, feature.sort_order), feature.sort_order)
    ordered_names = sorted(feature_names, key=lambda name: (feature_names[name], name))
    comparison_rows = []
    for name in ordered_names:
        cells = []
        for material in materials:
            matches = [feature for feature in material.features.all() if feature.name == name]
            cells.append(matches)
        comparison_rows.append({"name": name, "cells": cells})
    return render(request, "materials/material_compare.html", {
        "materials": materials,
        "comparison_rows": comparison_rows,
        "error": "",
    })


@login_required
def download_latest_excel(request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    index = workbook.active
    index.title = "分类目录"
    index.append(["材料特性分类目录"])
    index.append([])
    index.append(["分类名称", "材料数量"])
    categories = list(MaterialCategory.objects.annotate(material_count=Count("materials")))
    for category in categories:
        index.append([category.name, category.material_count])
    index.freeze_panes = "A4"
    index.column_dimensions["A"].width = 32
    index.column_dimensions["B"].width = 16

    for category in categories:
        sheet = workbook.create_sheet(title=category.name[:31])
        materials = list(
            Material.objects.filter(category=category)
            .prefetch_related("features")
            .order_by("name")
        )
        feature_names = {}
        for material in materials:
            for feature in material.features.all():
                feature_names[feature.name] = min(
                    feature_names.get(feature.name, feature.sort_order), feature.sort_order
                )
        ordered_names = sorted(feature_names, key=lambda name: (feature_names[name], name))
        headers = ["材料名称"]
        for name in ordered_names:
            headers.extend((f"{name}｜实验条件", f"{name}｜实验结果或数据"))
        headers.extend(("特殊备注", "引用文献"))
        sheet.append(headers)
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        for material in materials:
            features_by_name = {}
            for feature in material.features.all():
                features_by_name.setdefault(feature.name, []).append(feature)
            row = [material.name]
            for name in ordered_names:
                matches = features_by_name.get(name, [])
                row.extend((
                    "\n".join(item.condition for item in matches if item.condition),
                    "\n".join(item.result for item in matches if item.result),
                ))
            row.extend((material.notes, material.references))
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="243650")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.column_dimensions["A"].width = 38
        for column in range(2, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 42

    for cell in index[1]:
        cell.font = Font(bold=True, size=14, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243650")
    for cell in index[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3779E8")

    output = BytesIO()
    workbook.save(output)
    filename = "材料特性-最新数据.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f"attachment; filename=materials-latest.xlsx; filename*=UTF-8''{quote(filename)}"
    )
    return response
